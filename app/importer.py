import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.services import CategoryService, ProductService

REQUIRED_COLUMNS = ("name", "sku", "barcode", "price", "cost", "stock", "reorder", "category")


def create_template(path: str | Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(list(REQUIRED_COLUMNS))
    sheet.append(["Water", "WAT-001", "", 10.50, 7.00, 25, 5, "Drinks"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:H2"
    widths = {"A": 24, "B": 16, "C": 18, "D": 12, "E": 12, "F": 10, "G": 10, "H": 18}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    workbook.save(path)


def import_products(path: str | Path) -> tuple[int, list[str]]:
    path = Path(path)
    rows = _read_rows(path)
    imported = 0
    errors = []
    for row_number, row in enumerate(rows, start=2):
        try:
            values = {column: row.get(column, "") for column in REQUIRED_COLUMNS}
            if not values["name"] or not values["category"]:
                raise ValueError("name and category are required")
            category = next((item for item in CategoryService.list() if item.name.lower() == values["category"].strip().lower()), None)
            if category is None:
                CategoryService.save(values["category"])
                category = next(item for item in CategoryService.list() if item.name.lower() == values["category"].strip().lower())
            ProductService.create(
                _text(values["name"]),
                _text(values["sku"]),
                _text(values["barcode"]),
                _decimal(values["price"], "price"),
                _decimal(values["cost"], "cost"),
                _integer(values["stock"], "stock"),
                _integer(values["reorder"], "reorder"),
                category.id,
            )
            imported += 1
        except Exception as error:
            errors.append(f"Row {row_number}: {error}")
    return imported, errors


def _read_rows(path: Path) -> list[dict[str, object]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.DictReader(file)
            return _validate_columns(reader.fieldnames or [], list(reader))
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            values = list(workbook.active.values)
        finally:
            workbook.close()
        if not values:
            raise ValueError("The workbook is empty")
        headers = [str(value or "").strip().lower() for value in values[0]]
        return _validate_columns(headers, [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)])
    raise ValueError("Choose an .xlsx, .xlsm, or .csv file")


def _validate_columns(headers: list[str], rows: list[dict[str, object]]) -> list[dict[str, object]]:
    normalized = [header.strip().lower() for header in headers]
    missing = [column for column in REQUIRED_COLUMNS if column not in normalized]
    if missing:
        raise ValueError(f"Missing columns: {', '.join(missing)}")
    return [{str(key).strip().lower(): value for key, value in row.items()} for row in rows]


def _decimal(value: object, field: str):
    from decimal import Decimal
    try:
        return Decimal(str(value or "0").strip())
    except Exception as error:
        raise ValueError(f"{field} must be a number") from error


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _integer(value: object, field: str) -> int:
    try:
        return int(float(str(value or "0").strip()))
    except Exception as error:
        raise ValueError(f"{field} must be a whole number") from error
